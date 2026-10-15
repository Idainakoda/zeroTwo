"""Text extraction for uploaded files, with OCR for pictures and scanned pages.

Supported: plain text (.txt .md .csv .json .log), PDF, Word (.docx) and images
(.png .jpg .jpeg .webp .bmp .tif .tiff). Text inside pictures - a photographed
Aadhaar card, a screenshot embedded in a PDF, a scanned page - is read with OCR
so the detectors see it too. Configuration and credential files (.env and the like)
are refused: they are not documents and should never be uploaded to an assistant.
"""
import io
import logging
import re
import threading
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from app import config

log = logging.getLogger(__name__)

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
ALLOWED_EXT = TEXT_EXT | IMAGE_EXT | {".pdf", ".docx"}

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 30
MAX_OCR_IMAGES = 25
MAX_CHARS = 60_000
MIN_IMAGE_SIDE = 64          # icons, bullets and logos below this carry no readable text
SCANNED_PAGE_CHARS = 25      # a PDF page with less text than this is treated as a scan
OCR_DPI = 200


class ExtractionError(ValueError):
    """The file cannot be accepted (type, size or corrupt content). The message is shown to the user."""


@dataclass
class Extraction:
    filename: str
    kind: str
    text: str = ""
    pages: int = 0
    ocr_images: int = 0
    ocr_ms: float = 0.0
    truncated: bool = False
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- OCR

class _Ocr:
    """Lazy RapidOCR (ONNX) instance: about 15 MB of models, CPU only, loaded on first use."""

    def __init__(self):
        self._engine = None
        self._lock = threading.Lock()
        self.error = ""

    @property
    def ready(self) -> bool:
        return self._engine is not None

    def load(self) -> None:
        if self._engine is not None or not config.USE_OCR:
            return
        try:
            from rapidocr_onnxruntime import RapidOCR
            # All 28 logical cores oversubscribe ONNX Runtime; a few threads are 4x faster here.
            self._engine = RapidOCR(intra_op_num_threads=config.OCR_THREADS)
            self.error = ""
        except Exception as e:  # missing package or model: uploads still work, without OCR
            self.error = f"{type(e).__name__}: {e}"
            log.warning("OCR unavailable: %s", self.error)

    def read(self, image: bytes) -> str:
        self.load()
        if self._engine is None:
            return ""
        with self._lock:
            result, _ = self._engine(image, use_cls=False)
        return _join_lines(result or [])


ocr = _Ocr()

_CAMEL = re.compile(r"(?<=[a-z])(?=[A-Z])")


def _join_lines(result: list) -> str:
    """Rebuilds reading order from OCR boxes: rows top to bottom, boxes left to right."""
    boxes = []
    for box, text, score in result:
        if float(score) < 0.5 or not text.strip():
            continue
        ys = [p[1] for p in box]
        xs = [p[0] for p in box]
        # OCR often drops the space between words ("RameshKumarDOB"); split at case changes.
        boxes.append((min(ys), max(ys), min(xs), _CAMEL.sub(" ", text.strip())))
    boxes.sort(key=lambda b: (b[0], b[2]))
    lines: list[list[tuple]] = []
    for b in boxes:
        centre = (b[0] + b[1]) / 2
        if lines:
            top, bottom = min(x[0] for x in lines[-1]), max(x[1] for x in lines[-1])
            if top <= centre <= bottom:
                lines[-1].append(b)
                continue
        lines.append([b])
    return "\n".join("  ".join(b[3] for b in sorted(line, key=lambda x: x[2])) for line in lines)


# ---------------------------------------------------------------- extractors

def _ocr_block(out: Extraction, image: bytes, label: str) -> str:
    if out.ocr_images >= MAX_OCR_IMAGES:
        if "ocr-limit" not in out.warnings:
            out.warnings.append("ocr-limit")
        return ""
    t = time.perf_counter()
    text = ocr.read(image)
    out.ocr_ms += (time.perf_counter() - t) * 1000
    out.ocr_images += 1
    return f"[Text read from {label}]\n{text}" if text.strip() else ""


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
            if "\x00" not in text:
                return text
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf(data: bytes, out: Extraction) -> list[str]:
    import pymupdf as fitz

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ExtractionError("This PDF could not be opened. It may be damaged.") from e
    if doc.needs_pass:
        raise ExtractionError("This PDF is password-protected. Remove the password and upload it again.")
    parts: list[str] = []
    out.pages = doc.page_count
    if doc.page_count > MAX_PAGES:
        out.warnings.append("page-limit")
    for number, page in enumerate(doc, start=1):
        if number > MAX_PAGES:
            break
        page_text = page.get_text("text").strip()
        if page_text:
            parts.append(f"[Page {number}]\n{page_text}")
        images = page.get_images(full=True)
        if len(page_text) < SCANNED_PAGE_CHARS and (images or not page_text):
            # A scanned page: read the whole rendered page instead of its pieces.
            pix = page.get_pixmap(dpi=OCR_DPI)
            block = _ocr_block(out, pix.tobytes("png"), f"scanned page {number}")
            if block:
                parts.append(block)
            continue
        for index, img in enumerate(images, start=1):
            try:
                pix = fitz.Pixmap(doc, img[0])
                if pix.width < MIN_IMAGE_SIDE or pix.height < MIN_IMAGE_SIDE:
                    continue
                if pix.n - pix.alpha >= 4:          # CMYK -> RGB
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                block = _ocr_block(out, pix.tobytes("png"), f"image {index} on page {number}")
            except Exception as e:
                log.info("skipped an unreadable PDF image: %s", e)
                continue
            if block:
                parts.append(block)
    return parts


def _docx(data: bytes, out: Extraction) -> list[str]:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError("This Word file could not be opened. Save it as .docx and try again.") from e
    parts: list[str] = []
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    if paragraphs:
        parts.append("\n".join(paragraphs))
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        rows = [r for r in rows if r.strip(" |")]
        if rows:
            parts.append("\n".join(rows))
    index = 0
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        index += 1
        if _image_too_small(blob):
            continue
        block = _ocr_block(out, blob, f"picture {index} in the document")
        if block:
            parts.append(block)
    out.pages = 1
    return parts


def _image_too_small(blob: bytes) -> bool:
    try:
        from PIL import Image
        with Image.open(io.BytesIO(blob)) as im:
            return im.width < MIN_IMAGE_SIDE or im.height < MIN_IMAGE_SIDE
    except Exception:
        return True                                   # not a picture OCR can read (e.g. EMF)


def _image(data: bytes, out: Extraction) -> list[str]:
    from PIL import Image, UnidentifiedImageError

    try:
        with Image.open(io.BytesIO(data)) as im:
            im.load()
            frames = getattr(im, "n_frames", 1)
            pngs = []
            for frame in range(min(frames, MAX_PAGES)):
                im.seek(frame)
                rgb = im.convert("RGB")
                buf = io.BytesIO()
                rgb.save(buf, "PNG")
                pngs.append(buf.getvalue())
    except (UnidentifiedImageError, OSError) as e:
        raise ExtractionError("This picture could not be opened. Use PNG or JPG.") from e
    out.pages = len(pngs)
    parts = []
    for n, png in enumerate(pngs, start=1):
        block = _ocr_block(out, png, "the picture" if len(pngs) == 1 else f"page {n} of the picture")
        if block:
            parts.append(block)
    return parts


def _check_signature(ext: str, data: bytes) -> None:
    if ext == ".pdf" and not data.lstrip()[:5].startswith(b"%PDF"):
        raise ExtractionError("This file is named .pdf but is not a PDF.")
    if ext == ".docx" and not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExtractionError("This file is named .docx but is not a Word document.")


def extract(filename: str, data: bytes) -> Extraction:
    name = PurePosixPath((filename or "upload").replace("\\", "/")).name
    ext = PurePosixPath(name.lower()).suffix
    if not data:
        raise ExtractionError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ExtractionError("The file is larger than 10 MB.")
    if ext not in ALLOWED_EXT:
        raise ExtractionError("This file type is not supported. Upload a PDF, Word (.docx), picture "
                              "(PNG, JPG) or text file.")
    _check_signature(ext, data)

    if ext in TEXT_EXT:
        out = Extraction(name, "text", pages=1)
        parts = [_decode_text(data)]
    elif ext == ".pdf":
        out = Extraction(name, "pdf")
        parts = _pdf(data, out)
    elif ext == ".docx":
        out = Extraction(name, "docx")
        parts = _docx(data, out)
    else:
        out = Extraction(name, "image")
        parts = _image(data, out)

    if out.ocr_images and not ocr.ready:
        out.warnings.append("ocr-unavailable")
    text = "\n\n".join(p.strip() for p in parts if p and p.strip())
    if len(text) > MAX_CHARS:
        text, out.truncated = text[:MAX_CHARS], True
    out.text = text
    out.ocr_ms = round(out.ocr_ms, 1)
    return out
